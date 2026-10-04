from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Union

from app.utils.text_utils import normalize_whitespace


class UnsupportedFileTypeError(ValueError):
    """Raised when a resume is not a supported PDF or DOCX file."""


def extract_text(file: Union[str, Path, BinaryIO]) -> str:
    """Extract text from a PDF or DOCX path/file-like object."""
    name = getattr(file, "name", str(file))
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(file)
    if suffix == ".docx":
        return _extract_docx(file)
    raise UnsupportedFileTypeError(f"Unsupported file type: {suffix or 'unknown'}")


def _read_bytes(file: Union[str, Path, BinaryIO]) -> bytes:
    if isinstance(file, (str, Path)):
        return Path(file).read_bytes()
    if hasattr(file, "getvalue"):
        return file.getvalue()
    position = file.tell() if hasattr(file, "tell") else None
    data = file.read()
    if position is not None and hasattr(file, "seek"):
        file.seek(position)
    return data


def _extract_pdf(file: Union[str, Path, BinaryIO]) -> str:
    import fitz

    document = fitz.open(stream=_read_bytes(file), filetype="pdf")
    try:
        return normalize_whitespace("\n".join(page.get_text("text") for page in document))
    finally:
        document.close()


def _extract_docx(file: Union[str, Path, BinaryIO]) -> str:
    from docx import Document

    document = Document(BytesIO(_read_bytes(file)))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return normalize_whitespace("\n".join(paragraphs))
